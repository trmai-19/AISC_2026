import hashlib
import io
import json
import math
import os
import sys
import zipfile

from contextlib import closing
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from itertools import zip_longest
from pathlib import PurePosixPath

import boto3
import psycopg

from botocore.config import Config
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from psycopg.types.json import Jsonb
from pypdf import PdfReader


PARSER_VERSION = "external-v1"


DDL = """
CREATE SCHEMA IF NOT EXISTS pipeline_ops;
CREATE SCHEMA IF NOT EXISTS staging;

CREATE TABLE IF NOT EXISTS pipeline_ops.external_loads (
    run_id TEXT PRIMARY KEY,
    manifest_key TEXT NOT NULL,
    manifest_sha256 TEXT NOT NULL,
    parser_version TEXT NOT NULL,
    status TEXT NOT NULL,
    expected_files INTEGER NOT NULL,
    updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
);

CREATE TABLE IF NOT EXISTS staging.external_files (
    run_id TEXT NOT NULL
        REFERENCES pipeline_ops.external_loads(run_id),
    raw_key TEXT NOT NULL,
    source_id TEXT NOT NULL,
    filename TEXT NOT NULL,
    metadata JSONB NOT NULL,
    block_count BIGINT NOT NULL DEFAULT 0,
    loaded_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
    PRIMARY KEY (run_id, raw_key)
);

CREATE TABLE IF NOT EXISTS staging.external_blocks (
    run_id TEXT NOT NULL,
    raw_key TEXT NOT NULL,
    member_name TEXT NOT NULL,
    section_name TEXT NOT NULL,
    block_no INTEGER NOT NULL,
    block_type TEXT NOT NULL,
    payload JSONB NOT NULL,
    PRIMARY KEY (
        run_id, raw_key, member_name,
        section_name, block_no
    ),
    FOREIGN KEY (run_id, raw_key)
        REFERENCES staging.external_files(run_id, raw_key)
);
"""


def read_object(s3, bucket, key, version_id=None):
    request = {
        "Bucket": bucket,
        "Key": key,
    }

    if version_id:
        request["VersionId"] = version_id

    response = s3.get_object(**request)

    with response["Body"] as body:
        return body.read()


def json_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()

    if isinstance(value, (Decimal, timedelta)):
        return str(value)

    if isinstance(value, float) and not math.isfinite(value):
        return str(value)

    return value


def excel_blocks(data, member_name=""):
    # Đọc hai bản: giá trị tính sẵn và công thức gốc.
    with (
        closing(
            load_workbook(
                io.BytesIO(data),
                read_only=True,
                data_only=True,
                keep_links=False,
            )
        ) as values_book,
        closing(
            load_workbook(
                io.BytesIO(data),
                read_only=True,
                data_only=False,
                keep_links=False,
            )
        ) as formula_book,
    ):
        for sheet in values_book:
            formulas_sheet = formula_book[sheet.title]

            rows = zip_longest(
                sheet.iter_rows(values_only=True),
                formulas_sheet.iter_rows(),
                fillvalue=(),
            )

            for row_no, (values, cells) in enumerate(
                rows, start=1
            ):
                formulas = {
                    get_column_letter(index): getattr(
                        cell.value, "text", str(cell.value)
                    )
                    for index, cell in enumerate(cells, start=1)
                    if cell.data_type == "f"
                }

                if (
                    all(value is None for value in values)
                    and not formulas
                ):
                    continue

                yield (
                    member_name,
                    sheet.title,
                    row_no,
                    "excel_row",
                    {
                        "values": [
                            json_value(value) for value in values
                        ],
                        "formulas": formulas,
                    },
                )


def blocks(data, filename):
    suffix = PurePosixPath(filename).suffix.lower()

    if suffix == ".xlsx":
        yield from excel_blocks(data)

    elif suffix == ".zip":
        # QILT có XLSX và ODS; chỉ nạp bản XLSX.
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = [
                item
                for item in archive.infolist()
                if not item.is_dir()
                and item.filename.lower().endswith(".xlsx")
            ]

            if not members:
                raise ValueError("ZIP khong co XLSX de nap")

            names = [item.filename for item in members]

            if len(names) != len(set(names)):
                raise ValueError("ZIP co ten XLSX bi trung")

            for member in members:
                yield from excel_blocks(
                    archive.read(member),
                    member.filename,
                )

    elif suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))

        if not reader.pages:
            raise ValueError("PDF khong co trang")

        for page_no, page in enumerate(reader.pages, start=1):
            if page.get("/Contents") is None:
                text = ""
            else:
                text = (
                    page.extract_text(
                        extraction_mode="layout"
                    )
                    or ""
                )

            yield (
                "",
                "pages",
                page_no,
                "pdf_page",
                {
                    "text": text,
                    "needs_ocr_review": not bool(text.strip()),
                },
            )

    elif suffix in {".html", ".htm"}:
        soup = BeautifulSoup(data, "html.parser")

        title = (
            soup.title.get_text(" ", strip=True)
            if soup.title
            else None
        )

        for node in soup.select(
            "script,style,noscript,template,"
            "nav,header,footer,aside,form,svg"
        ):
            node.decompose()

        root = (
            soup.find("main")
            or soup.find("article")
            or soup.select_one('[role="main"]')
            or soup.body
        )

        if root is None:
            raise ValueError("HTML khong co vung noi dung")

        text = root.get_text("\n", strip=True)

        if len(text) < 200:
            raise ValueError("Noi dung HTML qua ngan")

        yield (
            "",
            "content",
            1,
            "html_document",
            {"title": title, "text": text},
        )

    else:
        raise ValueError(
            f"Chua ho tro dinh dang: {filename}"
        )


def main():
    if len(sys.argv) != 2:
        raise SystemExit(
            "Usage: python -m loading.load_external "
            "external/_runs/.../manifest.json"
        )

    manifest_key = sys.argv[1]
    bucket = os.environ["RAW_BUCKET"]

    s3 = boto3.client(
        "s3",
        endpoint_url=os.environ["S3_ENDPOINT_URL"],
        config=Config(
            signature_version="s3v4",
            s3={"addressing_style": "path"},
        ),
    )

    raw_manifest = read_object(
        s3, bucket, manifest_key
    )

    manifest = json.loads(raw_manifest)
    files = manifest.get("files", [])

    if (
        manifest.get("status") != "SUCCESS"
        or len(files) != 22
    ):
        raise ValueError(
            "Can manifest SUCCESS voi du 22 file External"
        )

    if len({item["key"] for item in files}) != len(files):
        raise ValueError("Manifest co raw key bi trung")

    if any(
        not item["key"].startswith("external/")
        for item in files
    ):
        raise ValueError(
            "Manifest chua key ngoai external/"
        )

    run_id = manifest["run_id"]

    manifest_sha = hashlib.sha256(
        raw_manifest
    ).hexdigest()

    errors = []

    with psycopg.connect(
        os.environ["TRANSFORMED_DATABASE_URL"],
        autocommit=True,
    ) as conn:
        with conn.transaction():
            conn.execute(
                "SELECT pg_advisory_xact_lock(20261019)"
            )

            conn.execute(DDL)

            existing = conn.execute(
                """
                SELECT
                    manifest_key,
                    manifest_sha256,
                    parser_version,
                    status
                FROM pipeline_ops.external_loads
                WHERE run_id = %s
                """,
                (run_id,),
            ).fetchone()

            if existing:
                if existing[:3] != (
                    manifest_key,
                    manifest_sha,
                    PARSER_VERSION,
                ):
                    raise ValueError(
                        "Run da ton tai voi manifest "
                        "hoac parser khac"
                    )

                if existing[3] == "SUCCESS":
                    print(
                        f"SKIP | Snapshot da nap: {run_id}",
                        flush=True,
                    )
                    return

            else:
                conn.execute(
                    """
                    INSERT INTO pipeline_ops.external_loads (
                        run_id,
                        manifest_key,
                        manifest_sha256,
                        parser_version,
                        status,
                        expected_files
                    )
                    VALUES (%s, %s, %s, %s, 'RUNNING', %s)
                    """,
                    (
                        run_id,
                        manifest_key,
                        manifest_sha,
                        PARSER_VERSION,
                        len(files),
                    ),
                )

        # Mỗi file là một transaction riêng.
        # File lỗi được rollback; file đã xong được giữ lại.
        for item in files:
            try:
                with conn.transaction():
                    conn.execute(
                        "SELECT pg_advisory_xact_lock(20261019)"
                    )

                    existing = conn.execute(
                        """
                        SELECT block_count
                        FROM staging.external_files
                        WHERE run_id = %s AND raw_key = %s
                        """,
                        (run_id, item["key"]),
                    ).fetchone()

                    if existing:
                        print(
                            f"SKIP FILE | {item['filename']}",
                            flush=True,
                        )
                        continue

                    print(
                        f"READ | {item['filename']}",
                        flush=True,
                    )

                    data = read_object(
                        s3,
                        bucket,
                        item["key"],
                        item.get("version_id"),
                    )

                    checksum = hashlib.sha256(data).hexdigest()

                    if (
                        len(data) != item["size_bytes"]
                        or checksum != item["sha256"]
                    ):
                        raise ValueError(
                            "Kich thuoc hoac SHA-256 khong khop"
                        )

                    metadata = {
                        **item,
                        "parser_version": PARSER_VERSION,
                    }

                    if item["filename"].lower().endswith(".zip"):
                        with zipfile.ZipFile(
                            io.BytesIO(data)
                        ) as archive:
                            metadata["ignored_archive_members"] = [
                                name
                                for name in archive.namelist()
                                if not name.endswith("/")
                                and not name.lower().endswith(".xlsx")
                            ]

                    conn.execute(
                        """
                        INSERT INTO staging.external_files (
                            run_id,
                            raw_key,
                            source_id,
                            filename,
                            metadata
                        )
                        VALUES (%s, %s, %s, %s, %s)
                        """,
                        (
                            run_id,
                            item["key"],
                            item["source_id"],
                            item["filename"],
                            Jsonb(metadata),
                        ),
                    )

                    count = 0
                    empty_pages = 0

                    with conn.cursor() as cur:
                        with cur.copy(
                            """
                            COPY staging.external_blocks (
                                run_id,
                                raw_key,
                                member_name,
                                section_name,
                                block_no,
                                block_type,
                                payload
                            )
                            FROM STDIN
                            """
                        ) as copy:
                            for block in blocks(
                                data, item["filename"]
                            ):
                                (
                                    member,
                                    section,
                                    number,
                                    kind,
                                    payload,
                                ) = block

                                copy.write_row((
                                    run_id,
                                    item["key"],
                                    member,
                                    section,
                                    number,
                                    kind,
                                    Jsonb(payload),
                                ))

                                count += 1
                                empty_pages += int(
                                    payload.get(
                                        "needs_ocr_review", False
                                    )
                                )

                    if count == 0:
                        raise ValueError(
                            "Khong trich xuat duoc block nao"
                        )

                    conn.execute(
                        """
                        UPDATE staging.external_files
                        SET block_count = %s
                        WHERE run_id = %s AND raw_key = %s
                        """,
                        (count, run_id, item["key"]),
                    )

                print(
                    f"LOADED | {item['filename']} | "
                    f"{count:,} blocks | "
                    f"empty_pdf_pages={empty_pages}",
                    flush=True,
                )

            except Exception as exc:
                errors.append(item["filename"])

                print(
                    f"FAILED | {item['filename']} | {exc}",
                    flush=True,
                )

        with conn.transaction():
            conn.execute(
                "SELECT pg_advisory_xact_lock(20261019)"
            )

            loaded = conn.execute(
                """
                SELECT COUNT(*)
                FROM staging.external_files
                WHERE run_id = %s
                """,
                (run_id,),
            ).fetchone()[0]

            status = (
                "SUCCESS"
                if loaded == len(files) and not errors
                else "PARTIAL"
            )

            conn.execute(
                """
                UPDATE pipeline_ops.external_loads
                SET status = %s, updated_at = NOW()
                WHERE run_id = %s
                """,
                (status, run_id),
            )

    print(
        f"\n{status} | files={loaded}/{len(files)} | "
        f"run_id={run_id}",
        flush=True,
    )

    if status != "SUCCESS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()